"""Exportación de los conceptos por trimestre a un libro Excel (.xlsx).

Genera un libro con una hoja por trimestre (Trimestre 1..4) con el mismo
formato que el Excel de referencia del autónomo:

    A1: INGRESOS            D1: GASTOS
    A2: CONCEPTO  B2: IMPORTE SIN IVA  C2: IVA   D2: CONCEPTO  E2: IMPORTE
"""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MESES = ["Ene-Mar", "Abr-Jun", "Jul-Sep", "Oct-Dic"]

_HEADER_FILL = PatternFill("solid", fgColor="1E1B4B")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, color="1E1B4B", size=12)
_TOTAL_FONT = Font(bold=True)
_TOTAL_FILL = PatternFill("solid", fgColor="EEF2FF")
_THIN = Side(style="thin", color="D8DEE9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MONEY = '#,##0.00 "€"'


def _write_side(ws, col_concepto: int, col_importe: int, col_iva: int | None, rows, total: Decimal) -> None:
    """Escribe una columna (ingresos o gastos) a partir de la fila 3."""
    r = 3
    for row in rows:
        concepto, importe = row[0], row[1]
        iva = row[2] if len(row) > 2 else None
        ws.cell(row=r, column=col_concepto, value=concepto).border = _BORDER
        c_imp = ws.cell(row=r, column=col_importe, value=float(importe))
        c_imp.number_format = _MONEY
        c_imp.border = _BORDER
        if col_iva is not None:
            c_iva = ws.cell(row=r, column=col_iva)
            if iva is None:
                c_iva.value = "N/A"
            else:
                c_iva.value = float(iva) / 100
                c_iva.number_format = "0%"
            c_iva.border = _BORDER
        r += 1

    ws.cell(row=r, column=col_concepto, value="TOTAL").font = _TOTAL_FONT
    ws.cell(row=r, column=col_concepto).fill = _TOTAL_FILL
    ws.cell(row=r, column=col_concepto).border = _BORDER
    c_tot = ws.cell(row=r, column=col_importe, value=float(total))
    c_tot.font = _TOTAL_FONT
    c_tot.fill = _TOTAL_FILL
    c_tot.number_format = _MONEY
    c_tot.border = _BORDER
    if col_iva is not None:
        ws.cell(row=r, column=col_iva).fill = _TOTAL_FILL
        ws.cell(row=r, column=col_iva).border = _BORDER


def generar_excel_trimestres(year: int, trimestres: list[dict]) -> bytes:
    """Devuelve el contenido .xlsx de los conceptos de un año.

    ``trimestres`` es una lista de diccionarios con las claves ``num``,
    ``ingresos`` y ``gastos``; cada ingreso es ``(concepto, importe, iva)`` y
    cada gasto ``(concepto, importe)``.
    """
    wb = Workbook()
    wb.remove(wb.active)

    for q in trimestres:
        num = q["num"]
        ws = wb.create_sheet(title=f"Trimestre {num}")

        ws["A1"] = "INGRESOS"
        ws["A1"].font = _TITLE_FONT
        ws["D1"] = "GASTOS"
        ws["D1"].font = _TITLE_FONT

        headers = {
            "A2": "CONCEPTO",
            "B2": "IMPORTE SIN IVA",
            "C2": "IVA",
            "D2": "CONCEPTO",
            "E2": "IMPORTE",
        }
        for cell, text in headers.items():
            c = ws[cell]
            c.value = text
            c.font = _HEADER_FONT
            c.fill = _HEADER_FILL
            c.alignment = Alignment(horizontal="center")
            c.border = _BORDER

        ingresos = q["ingresos"]
        gastos = q["gastos"]
        total_ingresos = sum((i[1] for i in ingresos), Decimal("0"))
        total_gastos = sum((g[1] for g in gastos), Decimal("0"))

        _write_side(ws, 1, 2, 3, ingresos, total_ingresos)
        _write_side(ws, 4, 5, None, gastos, total_gastos)

        ws.column_dimensions["A"].width = 46
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 10
        ws.column_dimensions["D"].width = 46
        ws.column_dimensions["E"].width = 18
        ws.freeze_panes = "A3"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
