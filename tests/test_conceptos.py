"""Tests del libro de conceptos por trimestre (CRUD, export a Excel y loader).

Los tests no dependen de datos reales: el seed está desactivado (SEED_CONCEPTOS=false
en conftest) y cada test crea sus propios conceptos.
"""

import json
from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl

YEAR = date.today().year


async def test_listado_estructura(auth):
    client, headers = auth
    r = await client.get(f"/api/v1/conceptos?year={YEAR}", headers=headers)
    assert r.status_code == 200, r.text
    data = r.json()

    assert data["year"] == YEAR
    assert sorted(q["num"] for q in data["trimestres"]) == [1, 2, 3, 4]
    assert all(q["ingresos"] == [] and q["gastos"] == [] for q in data["trimestres"])
    assert float(data["total_neto"]) == 0.0


async def test_crud_conceptos(auth):
    client, headers = auth

    r = await client.post("/api/v1/conceptos", headers=headers, json={
        "year": YEAR, "trimestre": 3, "tipo": "ingreso",
        "concepto": "Curso verano", "importe": 500, "iva": 21, "orden": 0,
    })
    assert r.status_code == 201, r.text
    c = r.json()
    assert c["tipo"] == "ingreso"

    r = await client.get(f"/api/v1/conceptos/{c['id']}", headers=headers)
    assert r.status_code == 200
    assert r.json()["concepto"] == "Curso verano"

    r = await client.put(f"/api/v1/conceptos/{c['id']}", headers=headers, json={
        "year": YEAR, "trimestre": 3, "tipo": "ingreso",
        "concepto": "Curso verano (editado)", "importe": 600, "iva": None, "orden": 0,
    })
    assert r.status_code == 200
    assert r.json()["concepto"] == "Curso verano (editado)"
    assert r.json()["iva"] is None

    r = await client.get(f"/api/v1/conceptos?year={YEAR}", headers=headers)
    t3 = next(q for q in r.json()["trimestres"] if q["num"] == 3)
    assert [x["concepto"] for x in t3["ingresos"]] == ["Curso verano (editado)"]
    assert float(t3["total_ingresos"]) == 600.0

    r = await client.delete(f"/api/v1/conceptos/{c['id']}", headers=headers)
    assert r.status_code == 204
    r = await client.get(f"/api/v1/conceptos?year={YEAR}", headers=headers)
    t3 = next(q for q in r.json()["trimestres"] if q["num"] == 3)
    assert t3["ingresos"] == []


async def test_export_excel(auth):
    client, headers = auth
    await client.post("/api/v1/conceptos", headers=headers, json={
        "year": YEAR, "trimestre": 1, "tipo": "ingreso",
        "concepto": "Venta ejemplo", "importe": 1000, "iva": 21, "orden": 0,
    })
    await client.post("/api/v1/conceptos", headers=headers, json={
        "year": YEAR, "trimestre": 1, "tipo": "ingreso",
        "concepto": "Ingreso exento", "importe": 500, "iva": None, "orden": 1,
    })
    await client.post("/api/v1/conceptos", headers=headers, json={
        "year": YEAR, "trimestre": 1, "tipo": "gasto",
        "concepto": "Gasto ejemplo", "importe": 100, "iva": None, "orden": 0,
    })

    r = await client.get(f"/api/v1/conceptos/export?year={YEAR}", headers=headers)
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"  # es un .xlsx (zip)

    wb = openpyxl.load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Trimestre 1", "Trimestre 2", "Trimestre 3", "Trimestre 4"]
    ws = wb["Trimestre 1"]
    assert ws["A1"].value == "INGRESOS"
    assert ws["D1"].value == "GASTOS"
    assert ws["A2"].value == "CONCEPTO"
    assert ws["B2"].value == "IMPORTE SIN IVA"
    assert ws["C2"].value == "IVA"
    assert ws["D2"].value == "CONCEPTO"
    assert ws["E2"].value == "IMPORTE"
    assert ws["A3"].value == "Venta ejemplo"
    assert ws["C3"].value == 0.21  # IVA como fracción, igual que el Excel original
    assert ws["A4"].value == "Ingreso exento"
    assert ws["C4"].value == "N/A"
    assert ws["D3"].value == "Gasto ejemplo"


async def test_aislamiento_entre_usuarios(auth, client):
    _, headers_a = auth
    await client.post("/api/v1/conceptos", headers=headers_a, json={
        "year": YEAR, "trimestre": 1, "tipo": "ingreso",
        "concepto": "Solo de A", "importe": 999, "iva": None, "orden": 0,
    })

    r = await client.post("/api/v1/auth/register", json={
        "email": "otro-conceptos@test.es", "password": "clave-segura-123", "full_name": "Otro",
    })
    headers_b = {"Authorization": f"Bearer {r.json()['access_token']}"}

    r = await client.get(f"/api/v1/conceptos?year={YEAR}", headers=headers_b)
    assert r.status_code == 200
    t1 = next(q for q in r.json()["trimestres"] if q["num"] == 1)
    assert t1["ingresos"] == []


def test_loader_seed(tmp_path, monkeypatch):
    import app.services.seed_conceptos as sc

    f = tmp_path / "seed.json"
    f.write_text(json.dumps({
        "1": {"ingresos": [["X", 100, 21], ["Y", 50, None]], "gastos": [["Z", 5]]},
    }), encoding="utf-8")
    monkeypatch.setattr(sc, "_SEED_FILE", f)

    data = sc._load()
    assert data[1]["ingresos"][0] == ("X", Decimal("100"), Decimal("21"))
    assert data[1]["ingresos"][1] == ("Y", Decimal("50"), None)
    assert data[1]["gastos"][0] == ("Z", Decimal("5"))


def test_loader_seed_sin_fichero(tmp_path, monkeypatch):
    import app.services.seed_conceptos as sc

    monkeypatch.setattr(sc, "_SEED_FILE", tmp_path / "no-existe.json")
    assert sc._load() == {}
